"""Local DOCX and PDF exports for a validated meeting protocol."""

from __future__ import annotations

from datetime import date
from html import escape
from io import BytesIO
from pathlib import Path
from typing import Iterable

from app.analysis import ActionItem, MeetingProtocol
from app.jobs import TranscriptSegment, TranscriptionResult


class ProtocolExportError(RuntimeError):
    """The local host cannot create a requested protocol export."""


_DOCUMENT_FONT = "Verdana"
_PDF_FONT_NAME = "TuyinUnicode"
_PDF_FONT_CANDIDATES = (
    Path("/System/Library/Fonts/Supplemental/Verdana.ttf"),
    Path("/System/Library/Fonts/Supplemental/Arial Unicode.ttf"),
    Path("/Library/Fonts/Arial Unicode.ttf"),
    Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
    Path("/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf"),
)


def create_protocol_export(
    export_format: str,
    result: TranscriptionResult,
    protocol: MeetingProtocol,
    *,
    generated_on: date | None = None,
) -> bytes:
    """Create an in-memory export from the current locally saved result."""
    if export_format == "docx":
        return _create_docx(result, protocol, generated_on=generated_on or date.today())
    if export_format == "pdf":
        return _create_pdf(result, protocol, generated_on=generated_on or date.today())
    raise ProtocolExportError("Поддерживаются только форматы DOCX и PDF.")


def _create_docx(result: TranscriptionResult, protocol: MeetingProtocol, *, generated_on: date) -> bytes:
    try:
        from docx import Document
        from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Inches, Pt, RGBColor
    except ImportError as error:  # pragma: no cover - dependency validation happens at startup/tests.
        raise ProtocolExportError("Для экспорта DOCX не установлена библиотека python-docx.") from error

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    normal = document.styles["Normal"]
    normal.font.name = _DOCUMENT_FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), _DOCUMENT_FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), _DOCUMENT_FONT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), _DOCUMENT_FONT)
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    _remove_docx_title_rule(document, qn=qn)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(12)
    _add_docx_run(title, "Протокол совещания", size=Pt(26), bold=False, color=RGBColor(0, 0, 0), qn=qn)

    _add_docx_metadata(document, result, protocol, generated_on, qn=qn, size=Pt(11))
    document.add_heading("Расшифровка", level=1)
    _style_docx_heading(document.paragraphs[-1], qn=qn, size=Pt(18), color=RGBColor(0, 0, 0))
    for segment in result.segments:
        speaker = _speaker_name(segment)
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(6)
        _add_docx_run(paragraph, f"{_timestamp(segment.start_seconds)}–{_timestamp(segment.end_seconds)}  ", size=Pt(9), color=RGBColor(89, 89, 89), qn=qn)
        _add_docx_run(paragraph, speaker, size=Pt(11), bold=True, qn=qn)
        text = document.add_paragraph(segment.text)
        text.paragraph_format.space_after = Pt(8)

    document.add_page_break()
    document.add_heading("Саммари", level=1)
    _style_docx_heading(document.paragraphs[-1], qn=qn, size=Pt(18), color=RGBColor(0, 0, 0))
    summary = document.add_paragraph(protocol.summary)
    summary.paragraph_format.space_after = Pt(10)
    if protocol.key_points:
        document.add_heading("Ключевые вопросы", level=2)
        _style_docx_heading(document.paragraphs[-1], qn=qn, size=Pt(14), color=RGBColor(0, 0, 0))
        for point in protocol.key_points:
            bullet = document.add_paragraph(style="List Bullet")
            bullet.add_run(point)

    document.add_heading("Поручения", level=1)
    _style_docx_heading(document.paragraphs[-1], qn=qn, size=Pt(18), color=RGBColor(0, 0, 0))
    if protocol.action_items:
        table = document.add_table(rows=1, cols=4)
        table.style = "Table Grid"
        table.autofit = False
        for cell, value, width in zip(table.rows[0].cells, ("Поручение", "Ответственный", "Срок", "Статус"), (3.0, 1.35, 1.15, 1.0)):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_docx_cell_margins(cell)
            _set_docx_cell_shading(cell, "1F4E79")
            _set_docx_cell_borders(cell)
            _set_docx_cell_text(cell, value, qn=qn, bold=True, color=RGBColor(255, 255, 255))
        _repeat_docx_header(table.rows[0])

        for item_index, item in enumerate(protocol.action_items):
            for chunk_index, description in enumerate(_description_chunks(item.description)):
                row = table.add_row()
                row_cells = row.cells
                values = (description, _assignee(item) if chunk_index == 0 else "", _deadline(item) if chunk_index == 0 else "", _status(item) if chunk_index == 0 else "")
                for cell, value, width in zip(row_cells, values, (3.0, 1.35, 1.15, 1.0)):
                    cell.width = Inches(width)
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    _set_docx_cell_margins(cell)
                    _set_docx_cell_borders(cell)
                    if item_index % 2:
                        _set_docx_cell_shading(cell, "EAF2F8")
                    _set_docx_cell_text(cell, value, qn=qn)
    else:
        document.add_paragraph("Явно сформулированных поручений в транскрипте не найдено.")

    properties = document.core_properties
    properties.author = "Tuyin"
    properties.last_modified_by = "Tuyin"
    properties.title = "Протокол совещания"
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def _create_pdf(result: TranscriptionResult, protocol: MeetingProtocol, *, generated_on: date) -> bytes:
    try:
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_CENTER
        from reportlab.lib.pagesizes import letter
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import inch
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    except ImportError as error:  # pragma: no cover - dependency validation happens at startup/tests.
        raise ProtocolExportError("Для экспорта PDF не установлена библиотека ReportLab.") from error

    font_path = next((path for path in _PDF_FONT_CANDIDATES if path.is_file()), None)
    if font_path is None:
        raise ProtocolExportError("Не найден локальный шрифт с поддержкой кириллицы и казахских символов.")
    if _PDF_FONT_NAME not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(_PDF_FONT_NAME, str(font_path)))

    stream = BytesIO()
    document = SimpleDocTemplate(
        stream,
        pagesize=letter,
        leftMargin=inch,
        rightMargin=inch,
        topMargin=inch,
        bottomMargin=inch,
        title="Протокол совещания",
        author="Tuyin",
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("ProtocolTitle", parent=styles["Title"], fontName=_PDF_FONT_NAME, fontSize=24, leading=29, alignment=TA_CENTER, textColor=colors.black, spaceAfter=14)
    metadata_style = ParagraphStyle("ProtocolMetadata", parent=styles["BodyText"], fontName=_PDF_FONT_NAME, fontSize=11, leading=15, textColor=colors.black, spaceAfter=5)
    body_style = ParagraphStyle("ProtocolBody", parent=styles["BodyText"], fontName=_PDF_FONT_NAME, fontSize=11, leading=15, textColor=colors.black, spaceAfter=8)
    speaker_style = ParagraphStyle("ProtocolSpeaker", parent=body_style, fontSize=11, leading=15, spaceBefore=6, spaceAfter=2)
    heading_style = ParagraphStyle("ProtocolHeading", parent=styles["Heading1"], fontName=_PDF_FONT_NAME, fontSize=18, leading=22, textColor=colors.black, spaceBefore=8, spaceAfter=10)
    subheading_style = ParagraphStyle("ProtocolSubheading", parent=styles["Heading2"], fontName=_PDF_FONT_NAME, fontSize=14, leading=18, textColor=colors.black, spaceBefore=8, spaceAfter=6)
    cell_style = ParagraphStyle("ProtocolCell", parent=body_style, fontSize=9, leading=12, spaceAfter=0)
    header_cell_style = ParagraphStyle("ProtocolHeaderCell", parent=cell_style, fontSize=9, leading=11, textColor=colors.white)

    story = [
        Paragraph("Протокол совещания", title_style),
        Paragraph(f"<b>Дата формирования:</b> {_export_date(generated_on)}", metadata_style),
        Paragraph(f"<b>Тема:</b> {_safe(protocol.title)}", metadata_style),
        Paragraph(f"<b>Участники:</b> {_safe(_participants(result))}", metadata_style),
        Spacer(1, 8),
        Paragraph("Расшифровка", heading_style),
    ]
    for segment in result.segments:
        story.append(Paragraph(f"<font size=9 color='#595959'>{_timestamp(segment.start_seconds)}–{_timestamp(segment.end_seconds)}</font>  <b>{_safe(_speaker_name(segment))}</b>", speaker_style))
        story.append(Paragraph(_safe(segment.text), body_style))

    story.extend([PageBreak(), Paragraph("Саммари", heading_style), Paragraph(_safe(protocol.summary), body_style)])
    if protocol.key_points:
        story.append(Paragraph("Ключевые вопросы", subheading_style))
        for point in protocol.key_points:
            story.append(Paragraph(_safe(point), body_style, bulletText="•"))

    story.append(Paragraph("Поручения", heading_style))
    if protocol.action_items:
        rows = [[Paragraph("Поручение", header_cell_style), Paragraph("Ответственный", header_cell_style), Paragraph("Срок", header_cell_style), Paragraph("Статус", header_cell_style)]]
        for item in protocol.action_items:
            for chunk_index, description in enumerate(_description_chunks(item.description)):
                rows.append(
                    [
                        Paragraph(_safe(description), cell_style),
                        Paragraph(_safe(_assignee(item) if chunk_index == 0 else ""), cell_style),
                        Paragraph(_safe(_deadline(item) if chunk_index == 0 else ""), cell_style),
                        Paragraph(_safe(_status(item) if chunk_index == 0 else ""), cell_style),
                    ]
                )
        table = Table(rows, colWidths=(3.0 * inch, 1.35 * inch, 1.15 * inch, 1.0 * inch), repeatRows=1, hAlign="LEFT")
        style_commands = [
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E79")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D9D9D9")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 6),
            ("RIGHTPADDING", (0, 0), (-1, -1), 6),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ]
        for index in range(1, len(rows)):
            if index % 2 == 0:
                style_commands.append(("BACKGROUND", (0, index), (-1, index), colors.HexColor("#EAF2F8")))
        table.setStyle(TableStyle(style_commands))
        story.append(table)
    else:
        story.append(Paragraph("Явно сформулированных поручений в транскрипте не найдено.", body_style))

    def add_footer(canvas, _document) -> None:
        canvas.saveState()
        canvas.setFont(_PDF_FONT_NAME, 8)
        canvas.setFillColor(colors.HexColor("#595959"))
        canvas.drawString(inch, 0.62 * inch, "Сформировано локально на этом компьютере")
        canvas.drawRightString(letter[0] - inch, 0.62 * inch, f"Страница {canvas.getPageNumber()}")
        canvas.restoreState()

    document.build(story, onFirstPage=add_footer, onLaterPages=add_footer)
    return stream.getvalue()


def _add_docx_metadata(document, result: TranscriptionResult, protocol: MeetingProtocol, generated_on: date, *, qn, size) -> None:
    for label, value in (
        ("Дата формирования", _export_date(generated_on)),
        ("Тема", protocol.title),
        ("Участники", _participants(result)),
    ):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = size
        _add_docx_run(paragraph, f"{label}: ", size=size, bold=True, qn=qn)
        _add_docx_run(paragraph, value, size=size, qn=qn)


def _remove_docx_title_rule(document, *, qn) -> None:
    """Remove Word's theme border so the Title style remains plain and black."""
    properties = document.styles["Title"]._element.get_or_add_pPr()
    for border in properties.findall(qn("w:pBdr")):
        properties.remove(border)


def _add_docx_run(paragraph, text: str, *, size, qn, bold: bool = False, color=None):
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = _DOCUMENT_FONT
    run._element.rPr.rFonts.set(qn("w:ascii"), _DOCUMENT_FONT)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), _DOCUMENT_FONT)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), _DOCUMENT_FONT)
    run.font.size = size
    if color is not None:
        run.font.color.rgb = color
    return run


def _style_docx_heading(paragraph, *, qn, size, color) -> None:
    paragraph.paragraph_format.space_before = size
    paragraph.paragraph_format.space_after = size
    for run in paragraph.runs:
        run.font.name = _DOCUMENT_FONT
        run._element.rPr.rFonts.set(qn("w:ascii"), _DOCUMENT_FONT)
        run._element.rPr.rFonts.set(qn("w:hAnsi"), _DOCUMENT_FONT)
        run._element.rPr.rFonts.set(qn("w:eastAsia"), _DOCUMENT_FONT)
        run.font.size = size
        run.font.color.rgb = color


def _set_docx_cell_text(cell, value: str, *, qn, bold: bool = False, color=None) -> None:
    from docx.shared import Pt

    paragraph = cell.paragraphs[0]
    paragraph.alignment = 0
    paragraph.paragraph_format.space_after = 0
    _add_docx_run(paragraph, value, size=Pt(9), bold=bold, color=color, qn=qn)


def _set_docx_cell_shading(cell, fill: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def _set_docx_cell_borders(cell) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    properties = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:color"), "D9D9D9")
        borders.append(border)
    properties.append(borders)


def _set_docx_cell_margins(cell) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    properties = cell._tc.get_or_add_tcPr()
    margins = OxmlElement("w:tcMar")
    for side in ("top", "left", "bottom", "right"):
        margin = OxmlElement(f"w:{side}")
        margin.set(qn("w:w"), "90")
        margin.set(qn("w:type"), "dxa")
        margins.append(margin)
    properties.append(margins)


def _repeat_docx_header(row) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    properties = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    properties.append(header)


def _participants(result: TranscriptionResult) -> str:
    names = [speaker.display_name for speaker in result.speakers]
    if not names:
        names = list(dict.fromkeys(_speaker_name(segment) for segment in result.segments))
    return ", ".join(names) if names else "Не определены"


def _speaker_name(segment: TranscriptSegment) -> str:
    return segment.speaker_name or segment.speaker_id or "Спикер не определён"


def _timestamp(seconds: float) -> str:
    whole_seconds = max(0, round(seconds))
    hours, remainder = divmod(whole_seconds, 3_600)
    minutes, seconds = divmod(remainder, 60)
    return f"{hours:02}:{minutes:02}:{seconds:02}" if hours else f"{minutes:02}:{seconds:02}"


def _export_date(value: date) -> str:
    return value.strftime("%d.%m.%Y")


def _assignee(item: ActionItem) -> str:
    return item.assignee or "Не указан"


def _deadline(item: ActionItem) -> str:
    exact = _export_date(item.deadline) if item.deadline else None
    if item.deadline_text and exact:
        return f"{item.deadline_text} ({exact})"
    return item.deadline_text or exact or "Не указан"


def _status(item: ActionItem) -> str:
    return "Новое" if item.status == "new" else item.status


def _description_chunks(text: str, *, maximum_length: int = 600) -> Iterable[str]:
    words = text.split()
    chunks: list[str] = []
    current: list[str] = []
    current_length = 0
    for word in words:
        if len(word) > maximum_length:
            if current:
                chunks.append(" ".join(current))
                current = []
                current_length = 0
            chunks.extend(word[index:index + maximum_length] for index in range(0, len(word), maximum_length))
            continue
        addition = len(word) + (1 if current else 0)
        if current and current_length + addition > maximum_length:
            chunks.append(" ".join(current))
            current = [word]
            current_length = len(word)
        else:
            current.append(word)
            current_length += addition
    if current:
        chunks.append(" ".join(current))
    return chunks or [text]


def _safe(value: str) -> str:
    return escape(value).replace("\n", "<br/>")
