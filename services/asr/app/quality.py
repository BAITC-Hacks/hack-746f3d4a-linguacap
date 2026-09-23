"""Privacy-aware, repeatable comparison of local transcripts and DOCX references.

The evaluator deliberately keeps transcript excerpts out of its public Markdown
summary.  Short expected/observed examples live only in the ignored runtime
JSON report, where they can be used to correct the locally stored transcript.
"""

from __future__ import annotations

from collections import Counter
from dataclasses import asdict, dataclass
from difflib import SequenceMatcher
from pathlib import Path
import re
from typing import Any, Iterable, Mapping, Sequence


_WORD_RE = re.compile(r"[0-9a-zа-яёәғқңөұүһі]+", re.IGNORECASE)
_CYRILLIC_RE = re.compile(r"[а-яәғқңөұүһі]", re.IGNORECASE)
_KAZAKH_SPECIFIC_RE = re.compile(r"[әғқңөұүһі]", re.IGNORECASE)
_TRANSCRIPT_HEADINGS = ("расшифровк", "транскрип", "стенограмм")
_SUMMARY_HEADINGS = ("саммари", "итог", "ключев", "поручен", "решени")
_METADATA_PREFIXES = ("протокол", "тема", "дата", "участник", "формат", "время")


@dataclass(frozen=True)
class SegmentRange:
    segment_id: str
    start_seconds: float
    end_seconds: float
    token_start: int
    token_end: int


@dataclass(frozen=True)
class Diagnostic:
    error_type: str
    language: str
    expected_words: int
    observed_words: int
    segment_id: str | None
    start_seconds: float | None
    end_seconds: float | None
    expected_excerpt: str
    observed_excerpt: str


def normalized_tokens(text: str) -> list[str]:
    """Return lowercase comparison tokens without punctuation or casing noise."""
    return [token.replace("ё", "е") for token in _WORD_RE.findall(text.casefold())]


def token_language(token: str) -> str:
    """Classify only reliably distinguishable Kazakh tokens; shared Cyrillic is RU.*"""
    if _KAZAKH_SPECIFIC_RE.search(token):
        return "kk"
    if _CYRILLIC_RE.search(token):
        return "ru"
    return "other"


def reference_from_docx(path: Path) -> tuple[list[str], list[str]]:
    """Extract transcript and action text from a supplied protocol DOCX.

    Reference documents are inputs only: their text is not logged or copied into
    the public report.  The function accepts variations in the heading wording
    used in the project templates.
    """
    try:
        from docx import Document
    except ImportError as error:  # pragma: no cover - dependency is pinned for the service.
        raise RuntimeError("Для сравнения нужен установленный python-docx.") from error

    document = Document(path)
    paragraphs = [
        (paragraph.text.strip(), paragraph.style.name if paragraph.style is not None else "")
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    ]
    transcript_start = 0
    heading_indexes: list[int] = []
    for index, (paragraph, style_name) in enumerate(paragraphs):
        if style_name.casefold().startswith("heading"):
            heading_indexes.append(index)
        lowered = paragraph.casefold()
        if style_name.casefold().startswith("heading") and any(marker in lowered for marker in _TRANSCRIPT_HEADINGS):
            transcript_start = index + 1
            break
    else:
        # The supplied protocols use a generic first Heading 1 (for example,
        # "Ход заседания") rather than a fixed word such as "Расшифровка".
        if heading_indexes:
            transcript_start = heading_indexes[0] + 1

    transcript_end = len(paragraphs)
    for index in range(transcript_start, len(paragraphs)):
        paragraph, style_name = paragraphs[index]
        lowered = paragraph.casefold()
        if style_name.casefold().startswith("heading") and any(marker in lowered for marker in _SUMMARY_HEADINGS):
            transcript_end = index
            break

    transcript = [
        paragraph
        for paragraph, style_name in paragraphs[transcript_start:transcript_end]
        if not style_name.casefold().startswith("heading") and _is_transcript_paragraph(paragraph)
    ]
    actions: list[str] = []
    for table in document.tables:
        if not table.rows:
            continue
        header = " ".join(cell.text.casefold() for cell in table.rows[0].cells)
        if "поруч" not in header and "задач" not in header:
            continue
        for row in table.rows[1:]:
            cells = [cell.text.strip() for cell in row.cells]
            if cells and cells[0]:
                actions.append(cells[0])
    return transcript, actions


def compare_tokens(
    reference_text: str,
    actual_segments: Sequence[Mapping[str, Any]],
    *,
    example_limit_per_language: int = 8,
) -> dict[str, Any]:
    """Compare exact normalized tokens and return aggregate metrics plus local diagnostics."""
    expected = normalized_tokens(reference_text)
    actual, ranges = _actual_tokens_and_ranges(actual_segments)
    matcher = SequenceMatcher(a=expected, b=actual, autojunk=False)
    reference_by_language = Counter(token_language(token) for token in expected)
    exact_by_language: Counter[str] = Counter()
    errors_by_language: Counter[str] = Counter()
    error_operations: dict[str, Counter[str]] = {language: Counter() for language in ("ru", "kk", "mixed", "other")}
    diagnostics: list[Diagnostic] = []
    kept_examples: Counter[str] = Counter()

    for tag, i1, i2, j1, j2 in matcher.get_opcodes():
        if tag == "equal":
            exact_by_language.update(token_language(token) for token in expected[i1:i2])
            continue
        language = _error_language(expected[i1:i2], actual[j1:j2])
        errors_by_language[language] += max(i2 - i1, j2 - j1)
        error_operations[language][tag] += 1
        if kept_examples[language] >= example_limit_per_language:
            continue
        segment = _segment_for_actual_index(ranges, j1)
        diagnostics.append(
            Diagnostic(
                error_type=tag,
                language=language,
                expected_words=i2 - i1,
                observed_words=j2 - j1,
                segment_id=segment.segment_id if segment else None,
                start_seconds=segment.start_seconds if segment else None,
                end_seconds=segment.end_seconds if segment else None,
                expected_excerpt=" ".join(expected[i1:i2][:8]),
                observed_excerpt=" ".join(actual[j1:j2][:8]),
            )
        )
        kept_examples[language] += 1

    exact_matches = sum(exact_by_language.values())
    return {
        "reference_words": len(expected),
        "transcript_words": len(actual),
        "exact_matches": exact_matches,
        "exact_token_coverage": _ratio(exact_matches, len(expected)),
        "language_coverage": {
            language: {
                "reference_words": reference_by_language[language],
                "exact_matches": exact_by_language[language],
                "coverage": _ratio(exact_by_language[language], reference_by_language[language]),
            }
            for language in ("ru", "kk", "other")
        },
        "error_units": {language: errors_by_language[language] for language in ("ru", "kk", "mixed", "other")},
        "error_operations": {
            language: {operation: error_operations[language][operation] for operation in ("replace", "delete", "insert")}
            for language in ("ru", "kk", "mixed", "other")
        },
        "diagnostics": [asdict(diagnostic) for diagnostic in diagnostics],
    }


def compare_actions(reference_actions: Iterable[str], actual_actions: Iterable[str]) -> dict[str, Any]:
    """Compare protocol actions by normalized token overlap; wording remains local-only."""
    expected = [_action_tokens(action) for action in reference_actions if _action_tokens(action)]
    actual = [_action_tokens(action) for action in actual_actions if _action_tokens(action)]
    available = set(range(len(actual)))
    matches = 0
    similarities: list[float] = []
    for reference_tokens in expected:
        best_index: int | None = None
        best_score = 0.0
        for index in available:
            score = _jaccard(reference_tokens, actual[index])
            if score > best_score:
                best_index = index
                best_score = score
        if best_index is not None and best_score >= 0.45:
            available.remove(best_index)
            matches += 1
            similarities.append(best_score)
    return {
        "reference_actions": len(expected),
        "protocol_actions": len(actual),
        "matched_actions": matches,
        "action_recall": _ratio(matches, len(expected)),
        "mean_matched_similarity": round(sum(similarities) / len(similarities), 3) if similarities else 0.0,
    }


def sanitized_markdown_report(comparisons: Sequence[Mapping[str, Any]], generated_on: str) -> str:
    """Render a Git-safe quality report with metrics but no transcript excerpts."""
    rows = [
        "| Запись | Эталонных слов | Слов в транскрипте | Точное покрытие | RU* | KK-специфичные | Эталонных поручений | Найдено |",
        "| --- | ---: | ---: | ---: | ---: | ---: | ---: | ---: |",
    ]
    details: list[str] = []
    for comparison in comparisons:
        transcript = comparison["transcript"]
        actions = comparison["actions"]
        language = transcript["language_coverage"]
        rows.append(
            "| {label} | {reference} | {actual} | {coverage} | {ru} | {kk} | {reference_actions} | {matched_actions} |".format(
                label=comparison["label"],
                reference=transcript["reference_words"],
                actual=transcript["transcript_words"],
                coverage=_percent(transcript["exact_token_coverage"]),
                ru=_coverage_cell(language["ru"]),
                kk=_coverage_cell(language["kk"]),
                reference_actions=actions["reference_actions"],
                matched_actions=actions["matched_actions"],
            )
        )
        errors = transcript["error_units"]
        operations = transcript["error_operations"]
        kk_status = (
            f"KK-специфичные: {errors['kk']} единиц расхождения."
            if language["kk"]["reference_words"]
            else "KK-специфичные: не проверяются — в эталоне нет слов с уникальными казахскими буквами."
        )
        mixed_status = (
            f"смешанные переключения: {errors['mixed']}."
            if language["kk"]["reference_words"]
            else "смешанные переключения: не проверяются этим эталоном."
        )
        details.extend(
            (
                f"### {comparison['label']}",
                "",
                f"- RU: {errors['ru']} единиц расхождения ({_operation_summary(operations['ru'])}); {kk_status} {mixed_status} Прочие: {errors['other']}.",
                f"- Поручения: {actions['matched_actions']} из {actions['reference_actions']} сопоставлены по лексическому пересечению; у локального протокола {actions['protocol_actions']} поручений.",
                "",
            )
        )
    return "\n".join(
        (
            "# Сравнение тестовых записей с эталонными протоколами",
            "",
            f"Дата проверки: {generated_on}.",
            "",
            "## Результаты",
            "",
            *rows,
            "",
            "*RU — все кириллические слова без специфичных казахских букв; поэтому метрика RU включает часть казахских слов с общей кириллицей. KK-специфичные — слова с `ә ғ қ ң ө ұ ү һ і`.*",
            "",
            "## Журнал категорий",
            "",
            *details,
            "## Метод и ограничения",
            "",
            "Сравнение нормализует регистр и пунктуацию и измеряет точное совпадение слов с DOCX-эталоном. Это не сертифицированный WER: эталон может быть редакторским, не синхронизирован по времени и использовать перефразирование. Поручения считаются совпавшими при пересечении нормализованных слов и основ не менее 45%, поэтому результат служит очередью ручной проверки, а не доказательством семантической эквивалентности.",
            "",
            "В этом файле нет фрагментов аудио, транскриптов, имён участников или текстов поручений. Короткие пары «ожидалось/распознано» и отметки времени сохраняются только в игнорируемом локальном JSON-отчёте, чтобы их можно было исправить через вкладку «Транскрипт».",
            "",
        )
    )


def _is_transcript_paragraph(paragraph: str) -> bool:
    lowered = paragraph.casefold().strip()
    if not lowered or any(lowered.startswith(prefix) for prefix in _METADATA_PREFIXES):
        return False
    # Speaker-only labels are short and have neither a terminal punctuation mark nor a digit timestamp.
    tokens = normalized_tokens(paragraph)
    return not (len(tokens) <= 3 and not re.search(r"[.!?…:]|\d", paragraph))


def _actual_tokens_and_ranges(actual_segments: Sequence[Mapping[str, Any]]) -> tuple[list[str], list[SegmentRange]]:
    tokens: list[str] = []
    ranges: list[SegmentRange] = []
    for segment in actual_segments:
        start = len(tokens)
        segment_tokens = normalized_tokens(str(segment.get("text", "")))
        tokens.extend(segment_tokens)
        ranges.append(
            SegmentRange(
                segment_id=str(segment.get("id", "")),
                start_seconds=float(segment.get("start_seconds", 0)),
                end_seconds=float(segment.get("end_seconds", 0)),
                token_start=start,
                token_end=len(tokens),
            )
        )
    return tokens, ranges


def _segment_for_actual_index(ranges: Sequence[SegmentRange], index: int) -> SegmentRange | None:
    if not ranges:
        return None
    for segment in ranges:
        if segment.token_start <= index < segment.token_end:
            return segment
    return ranges[min(len(ranges) - 1, max(0, index))]


def _error_language(expected: Sequence[str], actual: Sequence[str]) -> str:
    languages = {token_language(token) for token in (*expected, *actual)} - {"other"}
    if languages == {"ru", "kk"}:
        return "mixed"
    if "kk" in languages:
        return "kk"
    if "ru" in languages:
        return "ru"
    return "other"


def _jaccard(left: Sequence[str], right: Sequence[str]) -> float:
    left_set, right_set = set(left), set(right)
    return len(left_set & right_set) / len(left_set | right_set) if left_set or right_set else 0.0


def _action_tokens(text: str) -> list[str]:
    """Trim long inflectional endings without claiming full Russian/Kazakh lemmatization."""
    return [token if len(token) <= 6 else token[:6] for token in normalized_tokens(text)]


def _ratio(numerator: int, denominator: int) -> float:
    return round(numerator / denominator, 4) if denominator else 0.0


def _percent(value: float) -> str:
    return f"{value * 100:.1f}%"


def _coverage_cell(metric: Mapping[str, Any]) -> str:
    return _percent(float(metric["coverage"])) if metric["reference_words"] else "н/д"


def _operation_summary(operations: Mapping[str, int]) -> str:
    return "замены: {replace}; пропуски: {delete}; вставки: {insert}".format(**operations)
