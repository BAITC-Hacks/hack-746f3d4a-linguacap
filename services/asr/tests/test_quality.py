from pathlib import Path

from docx import Document

from app.quality import compare_actions, compare_tokens, normalized_tokens, reference_from_docx


def test_normalization_keeps_russian_and_kazakh_words_without_punctuation_noise():
    assert normalized_tokens("Ёлка, Әлия! 2026-жыл") == ["елка", "әлия", "2026", "жыл"]


def test_comparison_reports_kk_and_mixed_language_diagnostics_with_segment_location():
    metrics = compare_tokens(
        "Привет Әлия бүгін есепті жібер.",
        [
            {
                "id": "chunk-0001",
                "start_seconds": 12,
                "end_seconds": 20,
                "text": "Привет Алия сегодня отчет отправь.",
            }
        ],
    )

    assert metrics["reference_words"] == 5
    assert metrics["language_coverage"]["kk"]["reference_words"] == 4
    assert metrics["error_units"]["mixed"] > 0
    assert metrics["error_operations"]["mixed"]["replace"] == 1
    assert metrics["diagnostics"][0]["segment_id"] == "chunk-0001"


def test_action_comparison_matches_only_semantically_similar_token_sets():
    comparison = compare_actions(
        ["Подготовить отчет по бюджету", "Согласовать график встречи"],
        ["Подготовить бюджетный отчет", "Заказать пропуск"],
    )

    assert comparison["reference_actions"] == 2
    assert comparison["matched_actions"] == 1


def test_reference_extraction_uses_heading_boundaries_not_words_inside_transcript(tmp_path: Path):
    document = Document()
    document.add_heading("Ход заседания", level=1)
    document.add_paragraph("Обсудили поручение и сроки без завершения раздела.")
    document.add_heading("Саммари", level=1)
    document.add_paragraph("Краткий итог.")
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Поручение"
    table.rows[0].cells[1].text = "Срок"
    table.rows[1].cells[0].text = "Подготовить отчет"
    table.rows[1].cells[1].text = "Завтра"
    path = tmp_path / "reference.docx"
    document.save(path)

    transcript, actions = reference_from_docx(path)

    assert transcript == ["Обсудили поручение и сроки без завершения раздела."]
    assert actions == ["Подготовить отчет"]
